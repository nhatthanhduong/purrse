from __future__ import annotations

import calendar
import io
from pathlib import Path

from .database import Database, MonthSummary
from .insights import generate_insights
from .money import format_k


def _normalized_image(path: Path) -> io.BytesIO:
    """Return a report-safe, orientation-correct PNG for any supported gallery image."""
    from PIL import Image as PILImage, ImageOps
    output = io.BytesIO()
    with PILImage.open(path) as source:
        source.seek(0)
        image = ImageOps.exif_transpose(source)
        if image.mode not in ("RGB", "RGBA"):
            image = image.convert("RGBA")
        image.thumbnail((2400, 2400))
        image.save(output, format="PNG", optimize=True)
    output.seek(0)
    return output


def _title(summary: MonthSummary) -> str:
    return f"{calendar.month_name[summary.month]} {summary.year} Monthly Financial Report"


def _rows(summary: MonthSummary) -> list[tuple[str, str]]:
    closing_bank = summary.actual_bank if summary.actual_bank is not None else summary.expected_bank
    closing_debt = summary.actual_debt if summary.actual_debt is not None else summary.expected_debt
    return [
        ("Opening bank balance", format_k(summary.opening_bank)),
        ("Income", format_k(summary.income)),
        ("Expenses", format_k(summary.expenses)),
        ("Savings (income - expenses)", format_k(summary.savings)),
        ("Monthly savings target", format_k(summary.target_savings)),
        ("Money borrowed", format_k(summary.borrowed)),
        ("Debt repayments", format_k(summary.repayments)),
        ("Closing bank balance", format_k(closing_bank)),
        ("Closing debt", format_k(closing_debt)),
        ("Bank reconciliation", format_k(summary.bank_adjustment)),
        ("Debt reconciliation", format_k(summary.debt_adjustment)),
    ]


def _chart_png(categories: list[tuple[str, int]]) -> bytes | None:
    if not categories:
        return None
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib import pyplot as plt
    names = [row[0] for row in categories[:8]][::-1]
    values = [row[1] / 1000 for row in categories[:8]][::-1]
    fig, ax = plt.subplots(figsize=(7, 3.5))
    ax.barh(names, values, color="#d79ac8")
    ax.set_xlabel("thousand VND")
    ax.set_title("Spending by category")
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    output = io.BytesIO()
    fig.savefig(output, format="png", dpi=150, bbox_inches="tight")
    plt.close(fig)
    return output.getvalue()


def export_docx(db: Database, month_id: int, destination: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, RGBColor

    summary = db.summary(month_id)
    transactions = db.list_transactions(month_id)
    categories = db.category_totals(month_id)
    insights = generate_insights(summary, db.previous_summary(summary.year, summary.month), categories)
    document = Document()
    heading = document.add_heading("Purrse", 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subheading = document.add_paragraph(_title(summary))
    subheading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if summary.status == "open":
        draft = document.add_paragraph("DRAFT — this month has not been closed")
        draft.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in draft.runs:
            run.font.color.rgb = RGBColor(190, 80, 100)
    document.add_paragraph("All amounts are shown in thousand VND.")

    document.add_heading("Summary", level=1)
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    table.rows[0].cells[0].text = "Measure"
    table.rows[0].cells[1].text = "Amount"
    for label, value in _rows(summary):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value

    document.add_heading("Miso’s monthly notes", level=1)
    for line in insights:
        document.add_paragraph(line, style="List Bullet")
    if summary.note:
        document.add_heading("Personal reflection", level=1)
        document.add_paragraph(summary.note)

    chart = _chart_png(categories)
    if chart:
        document.add_heading("Spending overview", level=1)
        document.add_picture(io.BytesIO(chart), width=Inches(6.2))

    document.add_heading("Entries", level=1)
    entries = document.add_table(rows=1, cols=5)
    entries.style = "Light Shading Accent 1"
    for index, value in enumerate(("Date", "Type", "Category", "Description", "Amount")):
        entries.rows[0].cells[index].text = value
    for transaction in transactions:
        cells = entries.add_row().cells
        values = (transaction["tx_date"], transaction["kind"].title(), transaction["category"], transaction["description"], format_k(transaction["amount_vnd"]))
        for index, value in enumerate(values):
            cells[index].text = str(value)

    monthly_images = db.month_attachments(month_id)
    if monthly_images:
        document.add_page_break()
        document.add_heading("Monthly evidence gallery", level=1)
        document.add_paragraph("Balance screenshots and other evidence for this month, in the selected order.")
        for index, item in enumerate(monthly_images, start=1):
            path = db.attachment_dir / item["stored_name"]
            document.add_paragraph(item["caption"] or f"Image {index} — {item['original_name']}")
            if not path.exists():
                document.add_paragraph("Image file is missing.")
                continue
            try:
                document.add_picture(_normalized_image(path), width=Inches(5.8))
            except Exception:
                document.add_paragraph("Image could not be rendered.")

    attachments = db.attachments_for_month(month_id)
    images = [a for a in attachments if a["mime_type"] in ("image/jpeg", "image/png")]
    files = [a for a in attachments if a not in images]
    if images:
        document.add_page_break()
        document.add_heading("Receipt images", level=1)
        for item in images:
            path = db.attachment_dir / item["stored_name"]
            if not path.exists():
                document.add_paragraph(f"Missing attachment: {item['original_name']}")
                continue
            document.add_paragraph(item["caption"] or f"{item['tx_date']} — {item['category']} — {item['original_name']}")
            try:
                document.add_picture(str(path), width=Inches(5.8))
            except Exception:
                document.add_paragraph("Image could not be rendered.")
    if files:
        document.add_heading("Other attachments", level=1)
        for item in files:
            document.add_paragraph(f"{item['tx_date']} — {item['original_name']}", style="List Bullet")
    document.save(destination)
    return Path(destination)


def export_pdf(db: Database, month_id: int, destination: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    summary = db.summary(month_id)
    transactions = db.list_transactions(month_id)
    categories = db.category_totals(month_id)
    insights = generate_insights(summary, db.previous_summary(summary.year, summary.month), categories)
    font_name = "Helvetica"
    try:
        import matplotlib
        font_path = Path(matplotlib.get_data_path()) / "fonts" / "ttf" / "DejaVuSans.ttf"
        font_bold_path = Path(matplotlib.get_data_path()) / "fonts" / "ttf" / "DejaVuSans-Bold.ttf"
        pdfmetrics.registerFont(TTFont("PurrseSans", str(font_path)))
        pdfmetrics.registerFont(TTFont("PurrseSans-Bold", str(font_bold_path)))
        pdfmetrics.registerFontFamily("PurrseSans", normal="PurrseSans", bold="PurrseSans-Bold")
        font_name = "PurrseSans"
    except Exception:
        pass
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = font_name
    styles.add(ParagraphStyle(name="Centered", parent=styles["Title"], alignment=TA_CENTER, textColor=colors.HexColor("#8d557e")))
    doc = SimpleDocTemplate(str(destination), pagesize=A4, rightMargin=1.5 * cm, leftMargin=1.5 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    story = [Paragraph("Purrse", styles["Centered"]), Paragraph(_title(summary), styles["Heading2"])]
    if summary.status == "open":
        story.append(Paragraph("DRAFT — this month has not been closed", styles["Heading3"]))
    story.extend([Paragraph("All amounts are shown in thousand VND.", styles["BodyText"]), Spacer(1, 8)])
    data = [["Measure", "Amount"], *[list(row) for row in _rows(summary)]]
    table = Table(data, colWidths=[11 * cm, 5 * cm], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#ead4e4")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#b99aaf")),
        ("FONTNAME", (0, 0), (-1, 0), "PurrseSans-Bold" if font_name == "PurrseSans" else "Helvetica-Bold"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#fbf5f9")]),
    ]))
    story.extend([table, Spacer(1, 14), Paragraph("Miso’s monthly notes", styles["Heading2"])])
    for line in insights:
        story.append(Paragraph(f"• {line}", styles["BodyText"]))
    if summary.note:
        story.extend([Paragraph("Personal reflection", styles["Heading2"]), Paragraph(summary.note.replace("\n", "<br/>"), styles["BodyText"])])
    chart = _chart_png(categories)
    if chart:
        story.extend([Spacer(1, 10), Image(io.BytesIO(chart), width=16 * cm, height=8 * cm)])
    story.append(Paragraph("Entries", styles["Heading2"]))
    entry_data = [["Date", "Type", "Category", "Description", "Amount"]]
    for tx in transactions:
        entry_data.append([tx["tx_date"], tx["kind"].title(), tx["category"], Paragraph(tx["description"] or "", styles["BodyText"]), format_k(tx["amount_vnd"])])
    entry_table = Table(entry_data, colWidths=[2.3 * cm, 2.2 * cm, 3.2 * cm, 6 * cm, 2.3 * cm], repeatRows=1)
    entry_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#ead4e4")),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#b99aaf")),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(entry_table)

    monthly_images = db.month_attachments(month_id)
    if monthly_images:
        story.extend([PageBreak(), Paragraph("Monthly evidence gallery", styles["Heading1"]), Paragraph("Balance screenshots and other evidence for this month, in the selected order.", styles["BodyText"])])
        from PIL import Image as PILImage
        for index, item in enumerate(monthly_images, start=1):
            path = db.attachment_dir / item["stored_name"]
            story.append(Paragraph(item["caption"] or f"Image {index} — {item['original_name']}", styles["BodyText"]))
            if path.exists():
                try:
                    normalized = _normalized_image(path)
                    with PILImage.open(normalized) as source:
                        width, height = source.size
                    normalized.seek(0)
                    ratio = min((16 * cm) / width, (18 * cm) / height)
                    story.append(Image(normalized, width=width * ratio, height=height * ratio))
                except Exception:
                    story.append(Paragraph("Image could not be rendered.", styles["BodyText"]))
            else:
                story.append(Paragraph("Image file is missing.", styles["BodyText"]))
            story.append(Spacer(1, 12))

    attachments = db.attachments_for_month(month_id)
    images = [a for a in attachments if a["mime_type"] in ("image/jpeg", "image/png")]
    other = [a for a in attachments if a not in images]
    if images:
        story.extend([PageBreak(), Paragraph("Receipt images", styles["Heading1"])])
        for item in images:
            path = db.attachment_dir / item["stored_name"]
            story.append(Paragraph(item["caption"] or f"{item['tx_date']} — {item['category']} — {item['original_name']}", styles["BodyText"]))
            if path.exists():
                try:
                    from PIL import Image as PILImage
                    with PILImage.open(path) as source:
                        width, height = source.size
                    ratio = min((16 * cm) / width, (18 * cm) / height)
                    story.append(Image(str(path), width=width * ratio, height=height * ratio))
                except Exception:
                    story.append(Paragraph("Image could not be rendered.", styles["BodyText"]))
            else:
                story.append(Paragraph("Attachment file is missing.", styles["BodyText"]))
            story.append(Spacer(1, 12))
    if other:
        story.append(Paragraph("Other attachments", styles["Heading2"]))
        for item in other:
            story.append(Paragraph(f"• {item['tx_date']} — {item['original_name']}", styles["BodyText"]))
    doc.build(story)
    return Path(destination)
